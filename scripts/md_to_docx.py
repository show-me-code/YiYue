"""把项目的 Markdown 交付文档渲染成 DOCX（正式交付件）。

用途：`docs/交付说明书.md` 是**源**，本脚本负责生成可分发/可打印的 `.docx`。
改完 Markdown 后重跑即可：

    python scripts/md_to_docx.py docs/交付说明书.md docs/交付说明书.docx

依赖：`python-docx`（仅本脚本需要，应用本身不依赖它）。

支持的 Markdown 子集（够本项目文档用）：H1/H2/H3 标题、段落、`-` 无序列表、
`1.` 有序列表、`>` 引用、`|` 表格（含表头分隔行）、``` 围栏代码块、`---` 分隔线，
以及行内 `**粗体**`、`` `代码` ``、`[文字](目标)`。
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CJK_FONT = "Microsoft YaHei"
MONO_FONT = "Consolas"
BODY_SIZE = Pt(10.5)
TABLE_SIZE = Pt(9)

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))", re.S)

# 行尾是中文/中文标点时，续行直接拼接（中文不该插空格）；否则补一个空格，
# 避免把两个拉丁词粘成一个。
_CJK_TAIL = re.compile(r"[\u3000-\u303f\u4e00-\u9fff\uff00-\uffef" r"」』）】》”’。，、；：！？]$")


def _join_wrapped(lines: list[str]) -> str:
    """把 Markdown 的硬换行拼回一个段落（源文档为便于阅读按 100 列断行）。"""
    out = ""
    for line in lines:
        if not out:
            out = line
            continue
        out += "" if _CJK_TAIL.search(out) else " "
        out += line
    return out


def _set_cjk(run, font: str = CJK_FONT) -> None:
    """中文字体必须同时设置 w:eastAsia，否则 Word 里中文不会用该字体。"""
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font)


def _shade(paragraph, fill: str = "F2F3F5") -> None:
    pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def _add_inline(paragraph, text: str, size=BODY_SIZE) -> None:
    """把 `**粗体**` / `` `代码` `` / `[文字](目标)` 渲染成对应 run。"""
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = MONO_FONT
            run.font.size = Pt(size.pt - 0.5)
            run.font.color.rgb = RGBColor(0xB0, 0x30, 0x60)
        elif part.startswith("[") and "](" in part:
            label, _, target = part[1:-1].partition("](")
            shown = label if (label == target or target.startswith("http")) else f"{label}（{target}）"
            run = paragraph.add_run(shown)
            if target.startswith("http"):
                run.font.color.rgb = RGBColor(0x1A, 0x55, 0x9B)
                run.underline = True
        else:
            run = paragraph.add_run(part)
        run.font.size = size
        _set_cjk(run)


def _add_table(document: Document, rows: list[list[str]]) -> None:
    header, *body = rows
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, cell_text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        _add_inline(cell.paragraphs[0], cell_text, size=TABLE_SIZE)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in body:
        cells = table.add_row().cells
        for i, cell_text in enumerate(row[:len(header)]):
            cells[i].text = ""
            _add_inline(cells[i].paragraphs[0], cell_text, size=TABLE_SIZE)
    document.add_paragraph()


def _add_code(document: Document, lines: list[str]) -> None:
    for line in lines:
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.left_indent = Cm(0.5)
        run = p.add_run(line if line else " ")
        run.font.name = MONO_FONT
        run.font.size = Pt(9)
        _set_cjk(run, MONO_FONT)
        _shade(p)
    document.add_paragraph()


def _add_footer(document: Document, title: str) -> None:
    footer = document.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"{title}  ·  第 ")
    _set_cjk(run)
    run.font.size = Pt(8)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    num_run = p.add_run()
    num_run._r.append(fld_begin)
    num_run._r.append(instr)
    num_run._r.append(fld_end)
    tail = p.add_run(" 页")
    _set_cjk(tail)
    tail.font.size = Pt(8)


def convert(md_path: Path, docx_path: Path) -> None:
    text = md_path.read_text("utf-8")
    document = Document()
    style = document.styles["Normal"]
    style.font.name = CJK_FONT
    style.font.size = BODY_SIZE
    style.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    for section in document.sections:
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)  # A4
        section.left_margin = section.right_margin = Cm(2.0)
        section.top_margin = section.bottom_margin = Cm(2.0)

    lines = text.splitlines()
    table_buf: list[list[str]] = []
    code_buf: list[str] = []
    # 待排块：(种类, 样式名或 None, 行缓冲)。Markdown 里段落与列表项都可能被硬换行，
    # 续行必须归属到同一个块，否则"逐个字段一行"的正文会碎成很多段。
    pending: tuple[str, str | None, list[str]] | None = None
    in_code = False
    title = "译阅 YiYue 交付说明书"

    def flush_table() -> None:
        nonlocal table_buf
        if table_buf:
            _add_table(document, table_buf)
            table_buf = []

    def flush_pending() -> None:
        nonlocal pending
        if pending is None:
            return
        kind, style, buf = pending
        pending = None
        if not buf:
            return
        text_out = _join_wrapped(buf)
        if style:
            p = document.add_paragraph(style=style)
        else:
            p = document.add_paragraph()
        _add_inline(p, text_out)

    def start_pending(kind: str, style: str | None, first: str) -> None:
        nonlocal pending
        flush_pending()
        pending = (kind, style, [first])

    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        if line.startswith("```"):
            if in_code:
                _add_code(document, code_buf)
                code_buf, in_code = [], False
            else:
                flush_table()
                flush_pending()
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(raw)
            i += 1
            continue
        if line.startswith("|") and line.endswith("|"):
            flush_pending()
            cells = [c.strip() for c in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                i += 1
                continue
            table_buf.append(cells)
            i += 1
            continue
        flush_table()
        if not line.strip():
            flush_pending()
            i += 1
            continue
        if line.startswith("### "):
            flush_pending()
            document.add_heading(line[4:].strip(), level=3)
        elif line.startswith("## "):
            flush_pending()
            document.add_heading(line[3:].strip(), level=1)
        elif line.startswith("# "):
            flush_pending()
            title = line[2:].strip()
            document.add_heading(title, level=0)
        elif re.match(r"^[-*] ", line):
            start_pending("list", "List Bullet", line[2:].strip())
        elif re.match(r"^\d+\. ", line):
            start_pending("list", "List Number",
                          re.sub(r"^\d+\. ", "", line))
        elif line.startswith("> "):
            start_pending("quote", "Intense Quote", line[2:].strip())
        elif line.strip() == "---":
            flush_pending()
        elif pending is not None and pending[0] in ("list", "quote"):
            pending[2].append(line.strip())      # 列表项/引用的续行
        else:
            if pending is None or pending[0] != "para":
                start_pending("para", None, line.strip())
            else:
                pending[2].append(line.strip())
        i += 1
    flush_table()
    flush_pending()
    if in_code and code_buf:
        _add_code(document, code_buf)

    _add_footer(document, title)
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(docx_path))
    print(f"已生成 {docx_path}（{docx_path.stat().st_size} 字节）")


if __name__ == "__main__":
    src = Path(sys.argv[1] if len(sys.argv) > 1 else "docs/交付说明书.md")
    dst = Path(sys.argv[2] if len(sys.argv) > 2 else src.with_suffix(".docx"))
    convert(src, dst)
